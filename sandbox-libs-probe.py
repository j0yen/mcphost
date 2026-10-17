import io
import os

import openpyxl
import pypdf
from PIL import Image


def make_pdf(pages):
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        page_id = len(objs) + 1
        kids.append(f"{page_id} 0 R")
        stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET"
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page_id + 1} 0 R"
            " /Resources << /Font << /F1 3 0 R >> >> >>"
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>"
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


def main(args):
    reader = pypdf.PdfReader(io.BytesIO(make_pdf(["Alpha page one", "Bravo page two", "Charlie page three"])))
    pdf_text = " | ".join(p.extract_text().strip() for p in reader.pages)

    book = openpyxl.Workbook()
    sheet = book.active
    for i in range(100):
        sheet.append([i, f"row {i}"])
    buf = io.BytesIO()
    book.save(buf)
    buf.seek(0)
    xlsx_rows = sum(1 for _ in openpyxl.load_workbook(buf).active.iter_rows())

    png = io.BytesIO()
    Image.frombytes("RGB", (600, 600), os.urandom(600 * 600 * 3)).save(png, "PNG")
    png_in_bytes = png.tell()
    png.seek(0)
    small = Image.open(png).resize((256, 256))
    return {"pdf_text": pdf_text, "xlsx_rows": xlsx_rows, "png_px": small.size[0], "png_in_bytes": png_in_bytes}
